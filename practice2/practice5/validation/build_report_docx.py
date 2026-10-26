"""Build the student-facing practice 5 report from REPORT.md."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "REPORT.md"
OUTPUT = ROOT / "Практичне_заняття_5_PC_Builder.docx"


def add_inline(paragraph, content: str) -> None:
    content = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", content)
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", content)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        else:
            paragraph.add_run(part)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_borders(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:color"), "D9D9D9")
        borders.append(edge)
    tc_pr.append(borders)


def set_cell_margin(cell, top=75, start=85, bottom=75, end=85) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for side, width in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = OxmlElement(f"w:{side}")
        element.set(qn("w:w"), str(width))
        element.set(qn("w:type"), "dxa")
        margins.append(element)
    tc_pr.append(margins)


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    columns = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=columns)
    table.autofit = False
    available = 7.18
    max_lens = [max(len(row[col]) if col < len(row) else 0 for row in rows) for col in range(columns)]
    weights = [max(7, min(length, 30)) for length in max_lens]
    width_sum = sum(weights)
    widths = [Inches(available * weight / width_sum) for weight in weights]
    for col_no, width in enumerate(widths):
        table.columns[col_no].width = width
    for row_no, data in enumerate(rows):
        for col_no in range(columns):
            cell = table.cell(row_no, col_no)
            cell.width = widths[col_no]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.06
            value = data[col_no] if col_no < len(data) else ""
            add_inline(paragraph, value)
            for run in paragraph.runs:
                run.font.name = "Aptos"
                run.font.size = Pt(8.2)
                run.font.color.rgb = RGBColor(0, 0, 0)
                if row_no == 0:
                    run.bold = True
            set_cell_borders(cell)
            set_cell_margin(cell)
            if row_no == 0:
                set_cell_shading(cell, "DFEAF2")
            elif row_no % 2 == 0:
                set_cell_shading(cell, "F6F8FA")
    set_repeat_header(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_code(doc: Document, lines: list[str]) -> None:
    paragraph = doc.add_paragraph()
    paragraph.style = "Code"
    paragraph.paragraph_format.keep_together = len(lines) <= 16
    for index, line in enumerate(lines):
        if index:
            paragraph.add_run().add_break()
        paragraph.add_run(line)


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.66)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.66)
    section.right_margin = Inches(0.66)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.3)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for name, size, before, after in (("Title", 17, 0, 11), ("Heading 1", 12.5, 14, 7), ("Heading 2", 11, 10, 5)):
        style = styles[name]
        style.font.name = "Aptos"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    if "Code" not in styles:
        from docx.enum.style import WD_STYLE_TYPE
        code = styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code = styles["Code"]
    code.font.name = "Consolas"
    code.font.size = Pt(8)
    code.font.color.rgb = RGBColor(0, 0, 0)
    code.paragraph_format.left_indent = Inches(0.2)
    code.paragraph_format.space_before = Pt(3)
    code.paragraph_format.space_after = Pt(8)
    code.paragraph_format.line_spacing = 1.0

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Практичне заняття 5 · PC Builder").font.size = Pt(8)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("```"):
            code_lines = []
            index += 1
            while index < len(lines) and not lines[index].startswith("```"):
                code_lines.append(lines[index])
                index += 1
            add_code(doc, code_lines)
            index += 1
            continue
        if line.startswith("|") and index + 1 < len(lines) and re.match(r"^\|[\s|:\-]+\|$", lines[index + 1].strip()):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"[:\- ]+", cell) for cell in cells):
                    rows.append(cells)
                index += 1
            add_table(doc, rows)
            continue
        if line.startswith("# "):
            paragraph = doc.add_paragraph(style="Title")
            add_inline(paragraph, line[2:])
        elif line.startswith("## "):
            paragraph = doc.add_paragraph(style="Heading 1")
            add_inline(paragraph, line[3:])
        elif line.startswith("### "):
            paragraph = doc.add_paragraph(style="Heading 2")
            add_inline(paragraph, line[4:])
        elif line.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            add_inline(paragraph, line[2:])
        elif re.match(r"^\d+\. ", line):
            paragraph = doc.add_paragraph(style="List Number")
            add_inline(paragraph, re.sub(r"^\d+\. ", "", line))
        else:
            paragraph = doc.add_paragraph()
            add_inline(paragraph, line)
        index += 1

    doc.core_properties.title = "Практичне заняття 5 Базові SPARQL запити"
    doc.core_properties.subject = "PC Builder ontology SPARQL queries and inference comparison"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
